#!/usr/bin/env python3
from pathlib import Path
from datetime import date
from pptx import Presentation


def add_title_slide(prs: Presentation, title: str, subtitle: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle is not None:
        slide.placeholders[1].text = subtitle


def add_bulleted_slide(
    prs: Presentation,
    title: str,
    items: list[dict],  # each: {"text": str, "subs": list[str]}
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    first = True
    for item in items:
        text = item.get("text", "").strip()
        subs = item.get("subs", []) or []

        if not text:
            continue

        if first:
            text_frame.text = text
            para = text_frame.paragraphs[0]
            para.level = 0
            first = False
        else:
            p = text_frame.add_paragraph()
            p.text = text
            p.level = 0

        for sub in subs:
            sp = text_frame.add_paragraph()
            sp.text = sub
            sp.level = 1


def build_presentation(output_path: Path) -> None:
    prs = Presentation()

    # Titre
    add_title_slide(
        prs,
        title="Design Thinking — %ED",
        subtitle=f"Approche centrée utilisateur • {date.today().strftime('%d/%m/%Y')}",
    )

    # Définition
    add_bulleted_slide(
        prs,
        title="Définition (introduction)",
        items=[
            {
                "text": (
                    "Une approche d'innovation centrée sur l'humain qui combine "
                    "empathie, idéation, prototypage et test."
                ),
                "subs": [],
            },
            {
                "text": (
                    "Comprendre en profondeur les besoins des utilisateurs pour concevoir des "
                    "solutions désirables, faisables et viables."
                ),
                "subs": [],
            },
            {
                "text": "Processus itératif, collaboratif et orienté action.",
                "subs": [],
            },
        ],
    )

    # Étapes du processus
    add_bulleted_slide(
        prs,
        title="Étapes suivies par le Design Thinking (processus)",
        items=[
            {
                "text": "Empathie",
                "subs": [
                    "Observer, interviewer, cartographier l'expérience",
                    "Identifier motivations, douleurs, attentes",
                ],
            },
            {
                "text": "Définition",
                "subs": [
                    "Synthétiser les insights",
                    "Reformuler le problème sous forme de point de vue",
                ],
            },
            {
                "text": "Idéation",
                "subs": [
                    "Générer un grand volume d'idées (divergence)",
                    "Sélectionner et cadrer (convergence)",
                ],
            },
            {
                "text": "Prototypage",
                "subs": [
                    "Rendre tangible rapidement (maquettes, storyboards, cliquables)",
                    "Tester le concept à moindre coût",
                ],
            },
            {
                "text": "Test",
                "subs": [
                    "Recueillir des retours utilisateurs",
                    "Mesurer, apprendre, ajuster",
                ],
            },
            {
                "text": "Itération",
                "subs": [
                    "Boucler sur les étapes selon les retours",
                    "Améliorer la solution en continu",
                ],
            },
        ],
    )

    # Pourquoi c'est puissant
    add_bulleted_slide(
        prs,
        title="Pourquoi le Design Thinking est-il puissant ?",
        items=[
            {
                "text": "Centré utilisateur",
                "subs": ["Réduit le risque de construire une solution inadéquate"],
            },
            {
                "text": "Réduction du risque",
                "subs": ["Prototypes précoces → économies de temps et de budget"],
            },
            {
                "text": "Collaboration interdisciplinaire",
                "subs": ["Aligne parties prenantes et facilite l'adhésion"],
            },
            {
                "text": "Innovation pragmatique",
                "subs": ["Équilibre désirabilité, faisabilité, viabilité"],
            },
            {
                "text": "Accélération de l'apprentissage",
                "subs": ["Boucles rapides de feedback et d'amélioration"],
            },
        ],
    )

    # Conclusion
    add_bulleted_slide(
        prs,
        title="Conclusion",
        items=[
            {
                "text": (
                    "Se focaliser sur l'utilisateur, expérimenter rapidement, apprendre en continu."
                ),
                "subs": [],
            },
            {
                "text": "Appliquez-le par étapes: un sprint, un prototype, un test.",
                "subs": [],
            },
            {
                "text": "Mesurez l'impact utilisateur et business pour guider les décisions.",
                "subs": [],
            },
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


if __name__ == "__main__":
    out = Path("/workspace/outputs/Design_Thinking_ED.pptx")
    build_presentation(out)
    print(f"Fichier généré: {out}")
