# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Pt


def add_bulleted_slide(prs: Presentation, title: str, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1]
    tf = body.text_frame
    # Réinitialise proprement le text_frame pour conserver un paragraphe vide
    tf.clear()
    # Ajoute les puces
    for idx, item in enumerate(bullets):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = str(item), 0
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        for run in p.runs:
            run.font.size = Pt(24)


def build_presentation() -> Presentation:
    prs = Presentation()

    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Design Thinking"
    subtitle = slide.placeholders[1]
    subtitle.text = "Présentation (40 diapositives)"

    slides_data = [
        {"title": "Objectifs de la présentation", "bullets": [
            "Comprendre le Design Thinking",
            "Maîtriser son processus et ses livrables",
            "Savoir quand et pourquoi l'utiliser"
        ]},
        {"title": "Définition du Design Thinking", "bullets": [
            "Approche d'innovation centrée sur l'utilisateur",
            "Combine empathie, idéation, prototypage et test",
            "Favorise l'itération rapide et l'apprentissage"
        ]},
        {"title": "Principes clés", "bullets": [
            "Centré utilisateur",
            "Co-création pluridisciplinaire",
            "Itération et expérimentation",
            "Prototypage tangible",
            "Storytelling et visualisation"
        ]},
        {"title": "Origines et évolution", "bullets": [
            "Popularisé par IDEO et d.school (Stanford)",
            "Adoption large en produit, service, process",
            "Cadre adaptable aux contextes variés"
        ]},
        {"title": "Valeur pour les organisations", "bullets": [
            "Réduction des risques d'échec",
            "Alignement des équipes autour des besoins réels",
            "Accélération de l'apprentissage marché",
            "Meilleure adoption par les utilisateurs"
        ]},
        {"title": "Comparaison avec approches classiques", "bullets": [
            "Linéraire vs itératif",
            "Spécifications lourdes vs prototypes rapides",
            "Hypothèses implicites vs validation terrain",
            "Décisions top-down vs insights utilisateurs"
        ]},
        {"title": "Quand l'utiliser ?", "bullets": [
            "Incertitude sur le besoin ou la solution",
            "Recherche d'opportunités d'innovation",
            "Nécessité d'aligner des parties prenantes",
            "Amélioration d'expérience utilisateur"
        ]},
        {"title": "Processus en un coup d'œil", "bullets": [
            "1) Empathie",
            "2) Définition",
            "3) Idéation",
            "4) Prototypage",
            "5) Test"
        ]},
        # Empathie
        {"title": "Empathie — Objectif", "bullets": [
            "Comprendre les utilisateurs et leur contexte",
            "Observer comportements et motivations",
            "Identifier douleurs et besoins latents"
        ]},
        {"title": "Empathie — Méthodes", "bullets": [
            "Interviews, observations, shadowing",
            "Parcours et cartes d'empathie",
            "Journal de bord, immersions",
            "Collecte d'insights et citations"
        ]},
        {"title": "Empathie — Livrables", "bullets": [
            "Personas et segments",
            "Carte d'empathie / parcours",
            "Insights clés priorisés"
        ]},
        # Définition
        {"title": "Définition — Objectif", "bullets": [
            "Reformuler le problème à résoudre",
            "Cadrer le champ et la valeur visée",
            "Aligner l'équipe sur une POV claire"
        ]},
        {"title": "Définition — Outils", "bullets": [
            "Point de vue (POV)",
            "How Might We (HMW)",
            "Critères de succès",
            "Carte des hypothèses"
        ]},
        {"title": "Définition — Livrables", "bullets": [
            "Énoncé de problème centré utilisateur",
            "Hypothèses à tester",
            "Indicateurs de réussite"
        ]},
        # Idéation
        {"title": "Idéation — Objectif", "bullets": [
            "Générer un large éventail d'idées",
            "Favoriser la divergence avant convergence",
            "Explorer sans autocensure"
        ]},
        {"title": "Idéation — Techniques", "bullets": [
            "Brainstorming guidé",
            "Crazy 8s",
            "SCAMPER",
            "Analogies et benchmarks",
            "Co-création avec utilisateurs"
        ]},
        {"title": "Idéation — Critères de sélection", "bullets": [
            "Désirabilité (utilisateur)",
            "Faisabilité (technique)",
            "Viabilité (business)"
        ]},
        # Prototypage
        {"title": "Prototypage — Objectif", "bullets": [
            "Donner forme rapide aux idées",
            "Tester des hypothèses clés",
            "Apprendre au moindre coût"
        ]},
        {"title": "Prototypage — Types de prototypes", "bullets": [
            "Papier / fil de fer",
            "Maquettes interactives",
            "Service blueprint",
            "Pretotypes et simulations"
        ]},
        {"title": "Prototypage — Bonnes pratiques", "bullets": [
            "Construire juste assez",
            "Rendre testable une hypothèse",
            "Itérer vite, documenter les apprentissages",
            "Impliquer l'équipe et parties prenantes"
        ]},
        # Test
        {"title": "Test — Objectif", "bullets": [
            "Valider ou invalider des hypothèses",
            "Observer l'usage réel",
            "Recueillir feedbacks honnêtes"
        ]},
        {"title": "Test — Métriques et feedback", "bullets": [
            "Taux de réussite des tâches",
            "Satisfaction perçue",
            "Temps et erreurs",
            "Verbatims et surprises"
        ]},
        {"title": "Test — Boucle d'itération", "bullets": [
            "Adapter le prototype",
            "Reformuler si nécessaire",
            "Décider: pivoter, persévérer ou arrêter"
        ]},
        # Transversal
        {"title": "Itération et apprentissage", "bullets": [
            "Boucles rapides entre phases",
            "Cadre non linéaire",
            "Mesure continue de la valeur"
        ]},
        {"title": "Rôles et responsabilités", "bullets": [
            "Sponsor, facilitateur, designers",
            "Experts métier et techniques",
            "Utilisateurs et clients",
            "Décideurs et gouvernance"
        ]},
        {"title": "Cadence et rituels", "bullets": [
            "Ateliers time-boxés",
            "Revue d'apprentissages",
            "Démos fréquentes",
            "Rétrospectives"
        ]},
        {"title": "Outils et supports", "bullets": [
            "Miro, Figma, Mural",
            "Canva, Notion, Slides",
            "Templates HMW, Personas",
            "Guides d'interview"
        ]},
        # Cas
        {"title": "Étude de cas — Produit digital", "bullets": [
            "Refonte d'onboarding",
            "Tests de prototypes cliquables",
            "+20% activation",
            "Décisions guidées par données"
        ]},
        {"title": "Étude de cas — Service", "bullets": [
            "Parcours client repensé",
            "Formation front-line",
            "Réduction temps d'attente",
            "Satisfaction en hausse"
        ]},
        {"title": "Étude de cas — Process interne", "bullets": [
            "Simplification de workflow",
            "Automatisations ciblées",
            "Moins d'erreurs",
            "Gain de productivité"
        ]},
        # Pourquoi puissant
        {"title": "Pourquoi le Design Thinking est puissant", "bullets": [
            "Cadre simple et actionnable",
            "Orienté impact utilisateur",
            "Révèle l'inattendu",
            "Apprentissage rapide"
        ]},
        {"title": "Créativité structurée", "bullets": [
            "Divergence puis convergence",
            "Contraintes utiles",
            "Visualisation des idées",
            "Culture d'exploration"
        ]},
        {"title": "Réduction du risque", "bullets": [
            "Tester tôt, échouer vite et pas cher",
            "Décisions fondées sur preuves",
            "Alignement sur des hypothèses clés",
            "Éviter les grands paris aveugles"
        ]},
        {"title": "Alignement et collaboration", "bullets": [
            "Langage commun",
            "Co-création transverse",
            "Implication des parties prenantes",
            "Clarté des décisions"
        ]},
        {"title": "Focalisation valeur et ROI", "bullets": [
            "Désirabilité avant faisabilité",
            "Mesure des résultats",
            "Priorisation par impact",
            "Meilleure adoption"
        ]},
        {"title": "Accélération time-to-market", "bullets": [
            "Prototypes au lieu de specs",
            "Apprentissages en continu",
            "Réutilisation d'actifs",
            "Découpage en incréments"
        ]},
        # Bonnes pratiques / Pièges
        {"title": "Bonnes pratiques", "bullets": [
            "Clarifier le problème et le succès attendu",
            "Impliquer tôt les utilisateurs",
            "Documenter chaque apprentissage",
            "Time-boxer les ateliers",
            "Rendre visibles les décisions"
        ]},
        {"title": "Pièges à éviter", "bullets": [
            "Sauter l'empathie",
            "Prototyper trop tard ou trop parfait",
            "Confondre opinion et évidence",
            "Oublier de mesurer",
            "S'arrêter après un seul test"
        ]},
        # Conclusion
        {"title": "Conclusion et prochaines étapes", "bullets": [
            "Le Design Thinking: puissant et pragmatique",
            "Choisir un premier défi à adresser",
            "Former une équipe transverse",
            "Planifier un sprint d'exploration",
            "Mesurer et partager les résultats"
        ]},
    ]

    # Build slides
    for data in slides_data:
        add_bulleted_slide(prs, data["title"], data["bullets"]) 

    # Validate count = 40
    assert len(prs.slides) == 40, f"Nombre de diapositives inattendu: {len(prs.slides)}"

    return prs


def main():
    prs = build_presentation()
    prs.save("/workspace/Design_Thinking.pptx")


if __name__ == "__main__":
    main()
